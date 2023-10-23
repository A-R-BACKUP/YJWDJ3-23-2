import openpyxl as xl, docx, os

고객정보입력파일 = './input/name_addr.xlsx'
템플릿파일 = './input/event-template.docx'
저장폴더 = os.path.join(os.path.dirname(__file__),'output','events') 

if not os.path.exists(저장폴더):
    os.mkdir(저장폴더)

def read_data():
    result = []
    sheet = xl.load_workbook(고객정보입력파일).active
    for row in sheet.iter_rows(min_row=2):
        v = [c.value for c in row]
        if v[0] is None: break
        result.append(v)
    return result

for person in read_data():
    name,zipnum,addr = person
    card ={
        '[[주소]]': '(우)'+zipnum+'|'+addr,
        '[[고객명]]':name
    }
    doc = docx.Document(템플릿파일)
    for p in doc.paragraphs:
        for k,v in card.items():
            if p.text.find(k) >= 0:
                p.text = p.text.replace(k,v)
                p.runs[0].font.bold=True
    저장파일 = os.path.join(저장폴더,name+'님.docx');
    print('saved: ',저장파일)                

    doc.save(저장파일)